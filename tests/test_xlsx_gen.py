"""Tests for billed summary Excel export."""
from io import BytesIO

from openpyxl import load_workbook

from src.utils.xlsx_gen import generate_billed_summary_xlsx, safe_export_filename_part


def test_safe_export_filename_part_sanitizes_slash() -> None:
    """Billing refs with slashes become underscores."""
    assert safe_export_filename_part("INV/2026/01") == "INV_2026_01"
    assert safe_export_filename_part("  ") == "batch"


def test_generate_billed_summary_xlsx_vendor_subtotals() -> None:
    """Workbook has detail rows, vendor subtotals, and grand total."""
    context = {
        "invoice_number": "INV/99",
        "billing_date": "2026-07-30",
        "items": [
            {
                "serial_number": "R-1",
                "plate_number": "ABC 123",
                "vendor_name": "Shell",
                "fuel_type": "Diesel",
                "receipt_number": "RCPT1",
                "po_reference": "PO-1",
                "actual_quantity": 10,
                "actual_unit_price": 50,
                "total_price": 500,
            },
            {
                "serial_number": "R-2",
                "plate_number": "XYZ 9",
                "vendor_name": "Petron",
                "fuel_type": "Unleaded",
                "receipt_number": "RCPT2",
                "po_reference": "PO-2",
                "actual_quantity": 20,
                "actual_unit_price": 60,
                "total_price": 1200,
            },
            {
                "serial_number": "R-3",
                "plate_number": "DEF 4",
                "vendor_name": "Shell",
                "fuel_type": "Diesel",
                "receipt_number": "RCPT3",
                "po_reference": "PO-3",
                "actual_quantity": 5,
                "actual_unit_price": 50,
                "total_price": 250,
            },
        ],
    }
    raw = generate_billed_summary_xlsx(context)
    assert raw[:2] == b"PK"  # zip/xlsx magic

    workbook = load_workbook(BytesIO(raw))
    sheet = workbook.active
    assert sheet["B2"].value == "INV/99"
    assert sheet["B4"].value == 3

    values = []
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=8, values_only=True):
        values.append(row)

    labels = [row[2] for row in values if row[2]]
    assert any(isinstance(label, str) and label.startswith("Subtotal - Shell") for label in labels)
    assert any(isinstance(label, str) and label.startswith("Subtotal - Petron") for label in labels)
    assert "Grand Total" in labels
    assert any(row[7] == 1950 for row in values if row[2] == "Grand Total")
