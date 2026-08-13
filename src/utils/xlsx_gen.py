"""Excel (.xlsx) export helpers for billing summaries."""
from __future__ import annotations

import re
from io import BytesIO
from typing import Any, Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


def safe_export_filename_part(value: str, fallback: str = "batch") -> str:
    """
    Sanitize a string for use in a download filename.

    Args:
        value: Raw billing reference or similar.
        fallback: Used when value is empty after sanitizing.

    Returns:
        Filesystem-safe token.
    """
    cleaned = re.sub(r"[^\w.\-]+", "_", (value or "").strip(), flags=re.UNICODE)
    cleaned = cleaned.strip("._")
    return cleaned or fallback


def _as_float(value: Any) -> float:
    """Coerce a numeric field to float; invalid values become 0."""
    try:
        if value is None or value == "":
            return 0.0
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _group_items_by_vendor(items: List[Dict[str, Any]]) -> List[Tuple[str, List[Dict[str, Any]]]]:
    """
    Group billed items by vendor name (A–Z); unknown vendors last.

    Args:
        items: Billed requisition dicts.

    Returns:
        List of (vendor_label, rows) tuples.
    """
    buckets: Dict[str, List[Dict[str, Any]]] = {}
    for item in items:
        name = (item.get("vendor_name") or "").strip() or "— Unknown vendor"
        buckets.setdefault(name, []).append(item)

    def sort_key(label: str) -> Tuple[int, str]:
        if label.startswith("—"):
            return (1, label.upper())
        return (0, label.upper())

    ordered: List[Tuple[str, List[Dict[str, Any]]]] = []
    for label in sorted(buckets.keys(), key=sort_key):
        rows = sorted(
            buckets[label],
            key=lambda row: str(row.get("po_reference") or row.get("serial_number") or ""),
        )
        ordered.append((label, rows))
    return ordered


def generate_billed_summary_xlsx(context: Dict[str, Any]) -> bytes:
    """
    Build an Excel workbook for a just-billed vendor invoice batch.

    One sheet with detail rows, vendor subtotal rows, and a grand total.
    Context keys: invoice_number, billing_date, items (list of requisition dicts).

    Args:
        context: Billed summary payload (same shape as PDF summary).

    Returns:
        .xlsx file bytes.
    """
    invoice_number = (context.get("invoice_number") or "").strip()
    billing_date = (context.get("billing_date") or "").strip()
    items: List[Dict[str, Any]] = list(context.get("items") or [])

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Billed Summary"

    bold = Font(bold=True)
    header_font = Font(bold=True)
    money_format = "#,##0.00"

    sheet["A1"] = "Billed PO Summary"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = "Billing Reference"
    sheet["B2"] = invoice_number or "—"
    sheet["A3"] = "Billing Date"
    sheet["B3"] = billing_date or "—"
    sheet["A4"] = "Item Count"
    sheet["B4"] = len(items)

    headers = [
        "PO No.",
        "Plate",
        "Vendor",
        "Fuel Type",
        "Receipt #",
        "Actual Qty (L)",
        "Actual Unit Price",
        "Amount",
    ]
    header_row = 6
    for col_index, title in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=col_index, value=title)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    row_index = header_row + 1
    grand_total = 0.0

    for vendor_name, vendor_items in _group_items_by_vendor(items):
        vendor_total = 0.0
        for item in vendor_items:
            amount = _as_float(item.get("total_price"))
            vendor_total += amount
            qty = item.get("actual_quantity")
            if qty is None:
                qty = item.get("quantity")
            unit_price = item.get("actual_unit_price")
            if unit_price is None:
                unit_price = item.get("unit_price")

            sheet.cell(row=row_index, column=1, value=item.get("po_reference") or "")
            sheet.cell(row=row_index, column=2, value=item.get("plate_number") or "")
            sheet.cell(row=row_index, column=3, value=vendor_name)
            sheet.cell(row=row_index, column=4, value=item.get("fuel_type") or "")
            sheet.cell(row=row_index, column=5, value=item.get("receipt_number") or "")
            qty_cell = sheet.cell(row=row_index, column=6, value=_as_float(qty) if qty is not None else None)
            price_cell = sheet.cell(
                row=row_index,
                column=7,
                value=_as_float(unit_price) if unit_price is not None else None,
            )
            amount_cell = sheet.cell(row=row_index, column=8, value=amount)
            if price_cell.value is not None:
                price_cell.number_format = money_format
            amount_cell.number_format = money_format
            if qty_cell.value is not None:
                qty_cell.number_format = "0.##"
            row_index += 1

        # Vendor subtotal row
        sheet.cell(row=row_index, column=1, value="").font = bold
        label_cell = sheet.cell(row=row_index, column=3, value=f"Subtotal - {vendor_name}")
        label_cell.font = bold
        subtotal_cell = sheet.cell(row=row_index, column=8, value=vendor_total)
        subtotal_cell.font = bold
        subtotal_cell.number_format = money_format
        grand_total += vendor_total
        row_index += 1

    row_index += 1
    sheet.cell(row=row_index, column=3, value="Grand Total").font = bold
    grand_cell = sheet.cell(row=row_index, column=8, value=grand_total)
    grand_cell.font = bold
    grand_cell.number_format = money_format

    widths = [12, 14, 28, 16, 14, 14, 16, 12]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def generate_cc_reconcile_summary_xlsx(context: Dict[str, Any]) -> bytes:
    """
    Build an Excel workbook for a just-reconciled credit card batch.

    Context keys: statement_reference, reconciled_date, items.

    Args:
        context: CC reconcile summary payload.

    Returns:
        .xlsx file bytes.
    """
    statement_reference = (context.get("statement_reference") or "").strip()
    report_number = (context.get("report_number") or "").strip()
    reconciled_date = (context.get("reconciled_date") or "").strip()
    items: List[Dict[str, Any]] = list(context.get("items") or [])

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "CC Reconciliation"

    bold = Font(bold=True)
    header_font = Font(bold=True)
    money_format = "#,##0.00"

    sheet["A1"] = "CC Reconciliation Summary"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = "Report No."
    sheet["B2"] = report_number or "—"
    sheet["A3"] = "Statement Reference"
    sheet["B3"] = statement_reference or "—"
    sheet["A4"] = "Reconciled Date"
    sheet["B4"] = reconciled_date or "—"
    sheet["A5"] = "Item Count"
    sheet["B5"] = len(items)

    headers = [
        "Request #",
        "Plate",
        "Vendor",
        "Fuel Type",
        "Receipt #",
        "Card Ref",
        "Qty (L)",
        "Unit Price",
        "Amount",
    ]
    header_row = 7
    for col_index, title in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=col_index, value=title)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    row_index = header_row + 1
    grand_total = 0.0

    # Sort CC rows by request # within each vendor group.
    for vendor_name, vendor_items in _group_items_by_vendor(items):
        vendor_items = sorted(
            vendor_items,
            key=lambda row: str(row.get("serial_number") or ""),
        )
        vendor_total = 0.0
        for item in vendor_items:
            amount = _as_float(item.get("total_price"))
            vendor_total += amount
            qty = item.get("actual_quantity")
            if qty is None:
                qty = item.get("quantity")
            unit_price = item.get("actual_unit_price")
            if unit_price is None:
                unit_price = item.get("unit_price")

            sheet.cell(row=row_index, column=1, value=item.get("serial_number") or "")
            sheet.cell(row=row_index, column=2, value=item.get("plate_number") or "")
            sheet.cell(row=row_index, column=3, value=vendor_name)
            sheet.cell(row=row_index, column=4, value=item.get("fuel_type") or "")
            sheet.cell(row=row_index, column=5, value=item.get("receipt_number") or "")
            sheet.cell(row=row_index, column=6, value=item.get("cc_card_reference") or "")
            qty_cell = sheet.cell(
                row=row_index,
                column=7,
                value=_as_float(qty) if qty is not None else None,
            )
            price_cell = sheet.cell(
                row=row_index,
                column=8,
                value=_as_float(unit_price) if unit_price is not None else None,
            )
            amount_cell = sheet.cell(row=row_index, column=9, value=amount)
            if price_cell.value is not None:
                price_cell.number_format = money_format
            amount_cell.number_format = money_format
            if qty_cell.value is not None:
                qty_cell.number_format = "0.##"
            row_index += 1

        label_cell = sheet.cell(row=row_index, column=3, value=f"Subtotal - {vendor_name}")
        label_cell.font = bold
        subtotal_cell = sheet.cell(row=row_index, column=9, value=vendor_total)
        subtotal_cell.font = bold
        subtotal_cell.number_format = money_format
        grand_total += vendor_total
        row_index += 1

    row_index += 1
    sheet.cell(row=row_index, column=3, value="Grand Total").font = bold
    grand_cell = sheet.cell(row=row_index, column=9, value=grand_total)
    grand_cell.font = bold
    grand_cell.number_format = money_format

    widths = [12, 14, 28, 16, 14, 14, 12, 12, 12]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
